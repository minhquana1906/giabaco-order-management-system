import io
import uuid
from datetime import UTC, datetime
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.ext.asyncio import AsyncSession

from app.features.campaigns.models import Campaign
from app.features.orders.models import Order
from tests.features.orders.test_orders import Setup, put

PIVOT = "/api/stats/campaign-pivot"


async def submit(
    s: Setup,
    client: AsyncClient,
    session: AsyncSession,
    *,
    gender: str | None,
    kit_qty: int = 1,
    kit_size: str = "3",
    jacket: tuple[str, int] | None = None,
    submitted_at: datetime | None = None,
    submit_it: bool = True,
    **overrides: Any,
) -> uuid.UUID:
    oid = await s.order(session, gender=gender, **overrides)
    body = s.lines_body(
        kit_lines=[{"kit_id": s.kit_id, "qty": kit_qty, "size_label": kit_size}],
        item_lines=[{"product_id": s.jacket["id"], "size_label": jacket[0], "qty": jacket[1]}]
        if jacket
        else [],
        submit=submit_it,
    )
    r = await put(client, oid, body)
    assert r.status_code == 200, r.text
    if submitted_at:
        await session.execute(
            update(Order).where(Order.id == oid).values(submitted_at=submitted_at)
        )
        await session.commit()
    return oid


def counts(male: int, female: int, unknown: int) -> dict[str, int]:
    return {"male": male, "female": female, "unknown": unknown, "total": male + female + unknown}


async def fixture(client: AsyncClient, session: AsyncSession) -> tuple[Setup, dict[str, uuid.UUID]]:
    s = await Setup(client).build()
    ids = {
        "female": await submit(
            s, client, session, gender="FEMALE", jacket=("4", 1),
            submitted_at=datetime(2026, 7, 3, 17, 30, tzinfo=UTC),
        ),
        "male": await submit(
            s, client, session, gender="MALE", kit_qty=2,
            submitted_at=datetime(2026, 7, 3, 3, 0, tzinfo=UTC),
        ),
        "unknown": await submit(
            s, client, session, gender=None, kit_size="2", jacket=("4", 2),
            submitted_at=datetime(2026, 7, 1, 3, 0, tzinfo=UTC),
        ),
        "draft": await submit(s, client, session, gender="MALE", submit_it=False),
        "deleted": await submit(s, client, session, gender="MALE"),
    }  # fmt: skip
    await session.execute(
        update(Order).where(Order.id == ids["deleted"]).values(deleted_at=datetime.now(UTC))
    )
    await session.commit()
    return s, ids


def block(body: dict[str, Any], name_id: str) -> dict[str, Any]:
    return next(b for b in body["blocks"] if b["id"] == name_id)


def row(b: dict[str, Any], size: str) -> dict[str, Any]:
    return next(r for r in b["rows"] if r["size_label"] == size)


async def get(client: AsyncClient, s: Setup, **params: str) -> dict[str, Any]:
    r = await client.get(PIVOT, params={"campaign_id": s.cid, **params})
    assert r.status_code == 200, r.text
    return r.json()  # type: ignore[no-any-return]


async def test_product_qty_by_size_and_gender(client: AsyncClient, session: AsyncSession) -> None:
    s, _ = await fixture(client, session)
    body = await get(client, s)

    assert body["sizes"] == ["1", "2", "3", "4", "5"]
    assert [b["id"] for b in body["blocks"]][0] == s.shirt["id"]
    shirt = block(body, s.shirt["id"])
    assert row(shirt, "3") == {
        "size_label": "3",
        "male": 2,
        "female": 1,
        "unknown": 0,
        "total": 3,
        "qty": counts(2, 1, 0),
        "unpriced_qty": counts(0, 0, 0),
    }
    assert row(shirt, "2")["unknown"] == 1
    assert shirt["totals"] == {
        "male": 2, "female": 1, "unknown": 1, "total": 4,
        "qty": counts(2, 1, 1), "unpriced_qty": counts(0, 0, 0),
    }  # fmt: skip
    assert shirt["orders"] == 3
    jacket = block(body, s.jacket["id"])
    assert row(jacket, "4") == {
        "size_label": "4",
        "male": 0,
        "female": 1,
        "unknown": 2,
        "total": 3,
        "qty": counts(0, 1, 2),
        "unpriced_qty": counts(0, 0, 0),
    }
    assert body["grand_total"] == sum(b["totals"]["total"] for b in body["blocks"]) == 11


async def test_orders_measure_counts_each_order_once(
    client: AsyncClient, session: AsyncSession
) -> None:
    s, _ = await fixture(client, session)
    body = await get(client, s, measure="orders")

    shirt = block(body, s.shirt["id"])
    assert row(shirt, "3")["male"] == 1
    assert shirt["totals"]["total"] == 3
    assert block(body, s.jacket["id"])["totals"]["total"] == 2
    kit_body = await get(client, s, measure="orders", group="kit")
    assert kit_body["grand_total"] == 3


async def test_amount_measure_uses_line_snapshots(
    client: AsyncClient, session: AsyncSession
) -> None:
    s, _ = await fixture(client, session)
    body = await get(client, s, measure="amount")
    assert block(body, s.jacket["id"])["totals"]["total"] == 3 * 300_000
    assert block(body, s.shirt["id"])["totals"]["total"] == 4 * 100_000


async def test_kit_group_counts_kit_qty_by_common_size(
    client: AsyncClient, session: AsyncSession
) -> None:
    s, _ = await fixture(client, session)
    body = await get(client, s, group="kit")

    assert len(body["blocks"]) == 1
    kit = block(body, s.kit_id)
    assert row(kit, "3") == {
        "size_label": "3",
        "male": 2,
        "female": 1,
        "unknown": 0,
        "total": 3,
        "qty": counts(2, 1, 0),
        "unpriced_qty": counts(0, 0, 0),
    }
    assert row(kit, "2")["unknown"] == 1
    assert kit["totals"]["total"] == body["grand_total"] == 4
    assert kit["orders"] == 3
    amount = await get(client, s, group="kit", measure="amount")
    assert row(block(amount, s.kit_id), "3")["male"] == 2 * (100_000 + 140_000)


async def test_kit_size_tie_prefers_order_default(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session, gender="MALE", default_size="3")
    body = s.lines_body(item_lines=[], submit=True)
    body["kit_lines"][0]["slots"] = [{"slot_id": s.slots[0], "size_label": "5"}]
    assert (await put(client, oid, body)).status_code == 200
    kit = block(await get(client, s, group="kit"), s.kit_id)
    assert row(kit, "3")["male"] == 1 and row(kit, "5")["total"] == 0


async def test_submitted_until_cuts_at_vietnam_day_end(
    client: AsyncClient, session: AsyncSession
) -> None:
    s, _ = await fixture(client, session)
    body = await get(client, s, submitted_until="2026-07-03")
    assert block(body, s.shirt["id"])["totals"] == {
        "male": 2, "female": 0, "unknown": 1, "total": 3,
        "qty": counts(2, 0, 1), "unpriced_qty": counts(0, 0, 0),
    }  # fmt: skip
    assert (await get(client, s, submitted_until="2026-07-04"))["grand_total"] == 11
    assert (await get(client, s, submitted_until="2026-06-30"))["grand_total"] == 0


async def test_filters_and_validation(client: AsyncClient, session: AsyncSession) -> None:
    s, _ = await fixture(client, session)
    only = await get(client, s, product_id=s.jacket["id"])
    assert [b["id"] for b in only["blocks"]] == [s.jacket["id"]]
    by_kit = await get(client, s, kit_id=s.kit_id)
    assert block(by_kit, s.jacket["id"])["totals"]["total"] == 0
    assert (
        await client.get(PIVOT, params={"campaign_id": s.cid, "measure": "x"})
    ).status_code == 422
    assert (await client.get(PIVOT, params={"campaign_id": str(uuid.uuid4())})).status_code == 404


async def test_overview(client: AsyncClient, session: AsyncSession) -> None:
    s, _ = await fixture(client, session)
    r = await client.get("/api/stats/campaign-overview", params={"campaign_id": s.cid})
    body = r.json()
    assert (body["orders"], body["items"], body["unpriced_items"]) == (3, 11, 0)
    assert body["by_source"] == {"STAFF": 3}
    assert body["amount"] > 0
    until = await client.get(
        "/api/stats/campaign-overview",
        params={"campaign_id": s.cid, "submitted_until": "2026-06-30"},
    )
    assert until.json()["orders"] == 0


async def test_overview_counts_unpriced_items(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build(jacket_price=None)
    await submit(s, client, session, gender="MALE", jacket=("4", 2))
    body = (await client.get("/api/stats/campaign-overview", params={"campaign_id": s.cid})).json()
    assert (body["items"], body["unpriced_items"], body["amount"]) == (4, 2, 240_000)


async def test_overview_student_counts_and_priced_items(
    client: AsyncClient, session: AsyncSession, make_student: Any, make_class: Any
) -> None:
    s = await Setup(client).build(jacket_price=None)
    campaign = await session.get(Campaign, uuid.UUID(s.cid))
    assert campaign is not None
    klass = await make_class(1, "1A", school_id=campaign.school_id)
    kw = {"school_id": campaign.school_id, "class_id": klass.id}
    here = await make_student("A", **kw)
    second = await make_student("A2", **kw)
    deleted = await make_student("B", deleted=True, **kw)
    await make_student("C")
    await submit(s, client, session, gender="MALE", jacket=("4", 2), student_id=here.id)
    await submit(s, client, session, gender="MALE", student_id=second.id)
    await submit(s, client, session, gender="MALE", submit_it=False, student_id=deleted.id)
    await submit(s, client, session, gender="MALE")
    body = (await client.get("/api/stats/campaign-overview", params={"campaign_id": s.cid})).json()
    assert (body["students_total"], body["students_with_order"]) == (2, 2)
    assert (body["priced_items"], body["unpriced_items"]) == (6, 2)
    assert body["priced_items"] + body["unpriced_items"] == body["items"]


async def test_pivot_unpriced_qty(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build(jacket_price=None)
    await submit(s, client, session, gender="MALE", jacket=("4", 2))
    for group in ("product", "kit"):
        body = (
            await client.get(
                PIVOT, params={"campaign_id": s.cid, "measure": "amount", "group": group}
            )
        ).json()
        jacket = [b for b in body["blocks"] if b["id"] == s.jacket["id"]]
        if group == "product":
            assert jacket[0]["totals"]["unpriced_qty"] == counts(2, 0, 0)
            assert (
                row(jacket[0], "4")["unpriced_qty"] == row(jacket[0], "4")["qty"] == counts(2, 0, 0)
            )
            assert block(body, s.shirt["id"])["totals"]["unpriced_qty"] == counts(0, 0, 0)
        else:
            kit = block(body, s.kit_id)["totals"]
            assert kit["qty"] == counts(1, 0, 0) and kit["unpriced_qty"] == counts(0, 0, 0)


async def test_stats_and_export_require_session(anon_client: AsyncClient) -> None:
    cid = str(uuid.uuid4())
    for url in (
        f"{PIVOT}?campaign_id={cid}",
        f"/api/stats/campaign-overview?campaign_id={cid}",
        f"/api/export/campaign-orders.xlsx?campaign_id={cid}",
    ):
        assert (await anon_client.get(url)).status_code == 401


async def test_campaign_orders_export(client: AsyncClient, session: AsyncSession) -> None:
    s, _ = await fixture(client, session)
    r = await client.get("/api/export/campaign-orders.xlsx", params={"campaign_id": s.cid})
    assert r.status_code == 200
    workbook = load_workbook(io.BytesIO(r.content))
    assert workbook.sheetnames == ["Bảng kê đơn", "Ma trận"]
    headers, *rows = [list(x) for x in workbook["Bảng kê đơn"].iter_rows(values_only=True)]
    assert headers[:6] == ["Mã đơn", "Người nhận", "Phụ huynh", "Giới tính", "Lớp", "Cỡ chung"]
    assert headers[-3:] == ["Số bộ", "Thành tiền", "Ghi chú"]
    assert len(rows) == 3
    assert sum(row_[-3] for row_ in rows) == 4
    pivot_rows = [list(x) for x in workbook["Ma trận"].iter_rows(values_only=True)]
    kit_total = next(x for x in pivot_rows if x[0] == "Bộ" and x[2] == "Tổng")
    assert kit_total[3:] == [2, 1, 1, 4]


async def test_export_honours_submitted_until(client: AsyncClient, session: AsyncSession) -> None:
    s, _ = await fixture(client, session)
    r = await client.get(
        "/api/export/campaign-orders.xlsx",
        params={"campaign_id": s.cid, "submitted_until": "2026-07-03"},
    )
    assert r.status_code == 200
    workbook = load_workbook(io.BytesIO(r.content))
    assert len(list(workbook["Bảng kê đơn"].iter_rows(values_only=True))) - 1 == 2
    pivot_rows = [list(x) for x in workbook["Ma trận"].iter_rows(values_only=True)]
    kit_total = next(x for x in pivot_rows if x[0] == "Bộ" and x[2] == "Tổng")
    assert kit_total[3:] == [2, 0, 1, 3]


async def test_class_group_rows_are_sizes_per_class(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    await submit(s, client, session, gender="MALE", class_label="10A1")
    await submit(s, client, session, gender="FEMALE", class_label="2A1", jacket=("4", 1))
    await submit(s, client, session, gender="MALE", class_label="2A1", kit_qty=2)
    await submit(s, client, session, gender=None, kit_size="2")
    body = await get(client, s, group="class", product_id=s.shirt["id"])

    assert [b["name"] for b in body["blocks"]] == ["2A1", "10A1", "Chưa rõ lớp"]
    first = block(body, "2A1")
    assert row(first, "3") == {
        "size_label": "3",
        "male": 2,
        "female": 1,
        "unknown": 0,
        "total": 3,
        "qty": counts(2, 1, 0),
        "unpriced_qty": counts(0, 0, 0),
    }
    assert first["orders"] == 2
    assert row(block(body, ""), "2")["unknown"] == 1
    assert body["grand_total"] == sum(b["totals"]["total"] for b in body["blocks"]) == 5

    jackets = await get(client, s, group="class", product_id=s.jacket["id"])
    assert [b["name"] for b in jackets["blocks"]] == ["2A1"]
    assert jackets["grand_total"] == 1
