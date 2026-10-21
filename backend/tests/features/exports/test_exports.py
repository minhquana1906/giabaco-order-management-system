import csv
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from typing import Any

import pytest
from openpyxl import load_workbook

from app.features.campaigns.models import Campaign, CampaignItem
from app.features.exports.handler import (
    STUDENT_HEADERS,
    TRIAL_HEADERS,
    _sheet_title,
    to_csv,
    to_xlsx,
)
from app.features.orders.models import Order, OrderLine
from app.features.students.models import DataSource, Gender

URLS = [
    "/api/export/students.csv",
    "/api/export/students.xlsx",
    "/api/export/trials.csv",
    "/api/export/trials.xlsx",
    "/api/export/all.xlsx",
]


@pytest.fixture
async def seeded(make_student: Any) -> None:
    await make_student(
        "Nguyễn Minh Khang",
        [
            ("blazer", "2", "FITS", "3"),
            ("blazer", "4", "FITS", "3"),
            ("gile", "2", "TOO_TIGHT", "3"),
        ],
        note="Rộng vai, cần sửa",
    )
    await make_student("Lê Thị Ánh", [("blazer", "5", "TOO_LOOSE", None)], gender=Gender.FEMALE)
    await make_student("Đã Xóa", [("blazer", "3", "FITS", "3")], deleted=True)
    await make_student("Trần Văn Bình", [], deleted_trials=[("gile", "3", "FITS", "3")])


def read_csv(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def read_sheet(content: bytes, name: str) -> list[list[Any]]:
    sheet = load_workbook(io.BytesIO(content))[name]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


@pytest.mark.parametrize("url", URLS)
async def test_exports_require_session(anon_client: Any, url: str) -> None:
    assert (await anon_client.get(url)).status_code == 401


@pytest.mark.parametrize("url", URLS)
async def test_exports_are_attachments(client: Any, seeded: None, url: str) -> None:
    response = await client.get(url)

    assert response.status_code == 200
    filename = url.rsplit("/", 1)[1]
    assert response.headers["content-disposition"].startswith("attachment;")
    assert filename in response.headers["content-disposition"]


async def test_students_csv(client: Any, seeded: None) -> None:
    rows = read_csv((await client.get("/api/export/students.csv")).content)

    header, body = rows[0], rows[1:]
    assert header[: len(STUDENT_HEADERS)] == STUDENT_HEADERS
    assert header[len(STUDENT_HEADERS) :] == ["Tạm tính đơn (đ)", "Số lần thử"]
    assert not any("suy ra" in column for column in header)
    assert len(body) == 3
    names = {r[header.index("Họ và tên")] for r in body}
    assert names == {"Nguyễn Minh Khang", "Lê Thị Ánh", "Trần Văn Bình"}
    khang = next(r for r in body if r[3] == "Nguyễn Minh Khang")
    assert khang[header.index("Tuổi lúc đo")] == "8.0"
    assert khang[header.index("SĐT phụ huynh")] == "0912345678"
    assert khang[header.index("Số lần thử")] == "3"
    binh = next(r for r in body if r[3] == "Trần Văn Bình")
    assert binh[header.index("Số lần thử")] == "0"


async def test_trials_csv_is_raw_dataset(client: Any, seeded: None) -> None:
    rows = read_csv((await client.get("/api/export/trials.csv")).content)

    header, body = rows[0], rows[1:]
    assert header == TRIAL_HEADERS
    assert not any("suy ra" in column for column in header)
    assert len(body) == 4
    assert {r[header.index("Mức vừa")] for r in body} == {"FITS", "TOO_TIGHT", "TOO_LOOSE"}
    assert [r[header.index("Cỡ thử")] for r in body].count("3") == 0
    notes = [r[header.index("Ghi chú")] for r in body]
    assert "Rộng vai, cần sửa" in notes
    girl = next(r for r in body if r[header.index("Cỡ thử")] == "5")
    assert girl[header.index("Giới tính")] == "FEMALE"
    assert girl[header.index("Khối")] != ""
    assert girl[header.index("Cỡ gợi ý")] == ""
    boy = next(
        r for r in body if r[header.index("Cỡ thử")] == "2" and r[header.index("Sản phẩm")] != ""
    )
    assert boy[header.index("Cỡ gợi ý")] == "3"


async def test_xlsx_matches_csv(client: Any, seeded: None) -> None:
    students_csv = read_csv((await client.get("/api/export/students.csv")).content)
    students_xlsx = read_sheet((await client.get("/api/export/students.xlsx")).content, "Trường A")
    trials_csv = read_csv((await client.get("/api/export/trials.csv")).content)
    trials_xlsx = read_sheet((await client.get("/api/export/trials.xlsx")).content, "Lần thử")

    assert students_xlsx[0] == students_csv[0]
    assert len(students_xlsx) == len(students_csv) == 4
    assert {r[3] for r in students_xlsx[1:]} == {r[3] for r in students_csv[1:]}
    assert trials_xlsx[0] == trials_csv[0] == TRIAL_HEADERS
    assert len(trials_xlsx) == len(trials_csv) == 5


async def test_all_xlsx_has_both_sheets(client: Any, seeded: None) -> None:
    content = (await client.get("/api/export/all.xlsx")).content

    assert load_workbook(io.BytesIO(content)).sheetnames == ["Trường A", "Lần thử"]
    assert len(read_sheet(content, "Trường A")) == 4
    assert len(read_sheet(content, "Lần thử")) == 5


async def test_empty_database_exports_headers_only(client: Any) -> None:
    rows = read_csv((await client.get("/api/export/trials.csv")).content)

    assert rows == [TRIAL_HEADERS]


async def place_order(
    session: Any,
    campaign: Any,
    student: Any,
    lines: list[tuple[Any, str, int, int | None]],
    status: str = "SUBMITTED",
) -> None:
    order = Order(
        campaign_id=campaign.id,
        code=f"T{uuid.uuid4().hex[:10]}",
        student_id=student.id,
        recipient_name=student.full_name,
        source="PARENT",
        status=status,
        submitted_at=datetime.now(UTC) if status == "SUBMITTED" else None,
    )
    order.lines = [
        OrderLine(product_id=p.id, size_label=size, colour="ghi", qty=qty, unit_price=price)
        for p, size, qty, price in lines
    ]
    session.add(order)
    await session.commit()


async def test_students_export_has_data_source_and_order_columns(
    client: Any, make_student: Any, session: Any, products: Any, open_campaign: Any
) -> None:
    measured = await make_student("Nguyễn Minh Khang", [("blazer", "2", "FITS", "3")])
    await make_student("Lê Tự Khai", data_source=DataSource.PARENT_DECLARED)
    await make_student("Đã Xóa", deleted=True)
    campaign = await open_campaign(measured.school_id)
    await place_order(session, campaign, measured, [(products["blazer"], "3", 1, 500_000)])

    rows = read_csv((await client.get("/api/export/students.csv")).content)

    header, body = rows[0], rows[1:]
    assert {r[3] for r in body} == {"Nguyễn Minh Khang", "Lê Tự Khai"}
    by_name = {r[3]: r for r in body}
    for name, source in [
        ("Nguyễn Minh Khang", "ADMIN_MEASURED"),
        ("Lê Tự Khai", "PARENT_DECLARED"),
    ]:
        assert by_name[name][header.index("data_source")] == source
    assert not [h for h in header if "phụ huynh chọn" in h or "thời điểm chọn" in h]
    blazer = header.index("Áo Blazer (Vest) - cỡ đặt")
    assert by_name["Nguyễn Minh Khang"][blazer : blazer + 2] == ["3", "1"]
    assert by_name["Lê Tự Khai"][blazer : blazer + 2] == ["", ""]
    assert by_name["Lê Tự Khai"][header.index("Tạm tính đơn (đ)")] == ""
    assert by_name["Lê Tự Khai"][header.index("Số lần thử")] == "0"


async def test_trials_export_excludes_parent_declared_students(
    client: Any, make_student: Any
) -> None:
    await make_student("Nguyễn Minh Khang", [("blazer", "2", "FITS", "3")])
    await make_student(
        "Lê Tự Khai", [("blazer", "4", "FITS", "4")], data_source=DataSource.PARENT_DECLARED
    )

    csv_rows = read_csv((await client.get("/api/export/trials.csv")).content)
    xlsx_rows = read_sheet((await client.get("/api/export/trials.xlsx")).content, "Lần thử")
    all_rows = read_sheet((await client.get("/api/export/all.xlsx")).content, "Lần thử")

    assert len(csv_rows) == len(xlsx_rows) == len(all_rows) == 2
    assert csv_rows[1][csv_rows[0].index("Cỡ thử")] == "2"


FORMULA_NAMES = ["=1+1", "@SUM(1)", "+1", "-1", "\tcmd", "\rcmd"]
STUDENT_CELL_URLS = [
    "/api/export/students.csv",
    "/api/export/students.xlsx",
    "/api/export/all.xlsx",
]


def student_rows(url: str, content: bytes) -> list[list[Any]]:
    if url.endswith(".csv"):
        return read_csv(content)
    return [list(row) for row in load_workbook(io.BytesIO(content)).worksheets[0].values]


@pytest.mark.parametrize("url", STUDENT_CELL_URLS)
async def test_exports_neutralise_formula_cells_from_parent_names(
    client: Any,
    make_parent: Any,
    school_class: dict[str, Any],
    session: Any,
    url: str,
) -> None:
    parent = await make_parent()
    for i, name in enumerate(["=1+1", "@SUM(1)"]):
        r = await parent.post(
            "/api/parent/children",
            json={
                "school_id": school_class["school_id"],
                "class_id": school_class["class_id"],
                "full_name": name,
                "dob": f"2017-09-0{i + 1}",
                "gender": "FEMALE",
                "height_cm": 128.5,
                "weight_kg": 25.0,
            },
        )
        assert r.status_code == 201, r.text

    rows = student_rows(url, (await client.get(url)).content)

    names = {r[3] for r in rows[1:]}
    assert names == {"'=1+1", "'@SUM(1)"}


@pytest.mark.parametrize("url", URLS)
async def test_exports_neutralise_every_string_cell(
    client: Any, make_student: Any, url: str
) -> None:
    for i, name in enumerate(FORMULA_NAMES):
        await make_student(
            name,
            [("blazer", "2", "FITS", "3")],
            note=name,
            dob=date(2018, 3, i + 1),
        )

    content = (await client.get(url)).content
    if url.endswith(".csv"):
        sheets = [read_csv(content)]
    else:
        workbook = load_workbook(io.BytesIO(content))
        sheets = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in workbook.worksheets]
    cells = [c for rows in sheets for row in rows for c in row if isinstance(c, str)]

    assert cells
    assert not [c for c in cells if c.startswith(("=", "+", "-", "@", "\t", "\r"))]
    assert "'=1+1" in cells


def sheet_cells(url: str, content: bytes, sheet: str) -> list[Any]:
    rows = read_csv(content) if url.endswith(".csv") else read_sheet(content, sheet)
    return [cell for row in rows for cell in row]


@pytest.mark.parametrize("url", STUDENT_CELL_URLS)
async def test_students_exports_use_open_campaign_orders_and_drop_deleted_students(
    client: Any, make_student: Any, session: Any, products: Any, open_campaign: Any, url: str
) -> None:
    measured = await make_student("Nguyễn Minh Khang", [("blazer", "2", "FITS", "3")])
    declared = await make_student("Lê Tự Khai", data_source=DataSource.PARENT_DECLARED)
    gone = await make_student("Đã Xóa", [("blazer", "3", "FITS", "3")], deleted=True)
    campaign = await open_campaign(measured.school_id)
    await place_order(session, campaign, measured, [(products["blazer"], "3", 2, 500_000)])
    await place_order(
        session,
        campaign,
        declared,
        [(products["blazer"], "5", 1, 500_000), (products["gile"], "2", 1, None)],
    )
    await place_order(session, campaign, gone, [(products["blazer"], "1", 1, 500_000)])

    rows = student_rows(url, (await client.get(url)).content)

    header, body = rows[0], rows[1:]
    assert {r[3] for r in body} == {"Nguyễn Minh Khang", "Lê Tự Khai"}
    blazer = header.index("Áo Blazer (Vest) - cỡ đặt")
    gile = header.index("Áo Gile - cỡ đặt")
    total = header.index("Tạm tính đơn (đ)")
    by_name = {r[3]: [str(c) if c is not None else "" for c in r] for r in body}
    assert by_name["Nguyễn Minh Khang"][blazer : blazer + 2] == ["3", "2"]
    assert by_name["Nguyễn Minh Khang"][gile : gile + 2] == ["", ""]
    assert by_name["Nguyễn Minh Khang"][total] == "1000000"
    assert by_name["Lê Tự Khai"][blazer : blazer + 2] == ["5", "1"]
    assert by_name["Lê Tự Khai"][gile : gile + 2] == ["2", "1"]
    assert by_name["Lê Tự Khai"][total] == ""
    assert by_name["Lê Tự Khai"][header.index("data_source")] == "PARENT_DECLARED"
    assert "Đã Xóa" not in {c for r in rows for c in r}


TRIAL_URLS = ["/api/export/trials.csv", "/api/export/trials.xlsx", "/api/export/all.xlsx"]


@pytest.mark.parametrize("url", TRIAL_URLS)
async def test_trials_exports_never_carry_declared_students(
    client: Any, make_student: Any, url: str
) -> None:
    await make_student("Nguyễn Minh Khang", [("blazer", "2", "FITS", "3")], note="đo tại trường")
    await make_student(
        "Lê Tự Khai",
        [("gile", "6", "TOO_LOOSE", "6")],
        note="khai báo của phụ huynh",
        gender=Gender.FEMALE,
        data_source=DataSource.PARENT_DECLARED,
    )

    cells = sheet_cells(url, (await client.get(url)).content, "Lần thử")

    assert "đo tại trường" in cells
    for declared_marker in ("khai báo của phụ huynh", "Lê Tự Khai", "6", "TOO_LOOSE", "FEMALE"):
        assert declared_marker not in cells


CRAFTED_ROW = ["=SUM(A1)", "\tcmd", "\rcmd", "+1", "-1", "@a", "plain", "Ánh", 5, -3, 2.5, None]


def test_to_csv_prefixes_only_formula_triggering_text() -> None:
    rows = read_csv(to_csv(["h"], [CRAFTED_ROW]))

    assert rows[1] == [
        "'=SUM(A1)",
        "'\tcmd",
        "'\rcmd",
        "'+1",
        "'-1",
        "'@a",
        "plain",
        "Ánh",
        "5",
        "-3",
        "2.5",
        "",
    ]


def test_to_xlsx_prefixes_text_and_keeps_numbers_and_empty_cells() -> None:
    content = to_xlsx({"S": (["=h"], [CRAFTED_ROW])})

    sheet = load_workbook(io.BytesIO(content))["S"]
    header, row = list(sheet.iter_rows())
    assert header[0].value == "'=h" and header[0].data_type == "s"
    assert [c.value for c in row[:8]] == [
        "'=SUM(A1)",
        "'\tcmd",
        "'\ncmd",  # XML turns the stored CR into LF on read; the quote prefix is what matters
        "'+1",
        "'-1",
        "'@a",
        "plain",
        "Ánh",
    ]
    assert all(c.data_type == "s" for c in row[:8])
    assert [(c.value, c.data_type) for c in row[8:11]] == [(5, "n"), (-3, "n"), (2.5, "n")]
    assert len(row) == 11


async def _blazer_only_campaign(session: Any, products: Any, school_id: Any) -> None:
    campaign = Campaign(
        school_id=uuid.UUID(str(school_id)),
        name="Đợt blazer",
        school_year="2026-2027",
        opens_at=datetime.now(UTC) - timedelta(days=1),
        status="LOCKED",
    )
    campaign.items = [CampaignItem(product_id=products["blazer"].id, colour="ghi", sort=0)]
    session.add(campaign)
    await session.commit()


@pytest.fixture
async def two_schools(
    client: Any, make_student: Any, make_school: Any, make_class: Any, session: Any, products: Any
) -> dict[str, Any]:
    first = await make_student("Nguyễn Minh Khang")
    other = await make_school("Trường Q: [LQD]/Một", "LQD")
    other_class = await make_class(2, "2Q1", other.id)
    await make_student("Lê Thị Ánh", school_id=other.id, class_id=other_class.id)
    empty = await make_school("Trường Trống", "TT")
    await _blazer_only_campaign(session, products, other.id)
    return {"first": first.school_id, "other": other.id, "empty": empty.id}


async def test_students_xlsx_has_one_sheet_per_school_with_only_its_profile(
    client: Any, two_schools: dict[str, Any], session: Any, products: Any, open_campaign: Any
) -> None:
    await open_campaign(two_schools["first"])

    content = (await client.get("/api/export/students.xlsx")).content

    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Trường A", "Trường Q   LQD  Một", "Trường Trống"]
    tail = ["Tạm tính đơn (đ)", "Số lần thử"]
    first, other, empty = (list(ws.iter_rows(values_only=True)) for ws in workbook.worksheets)
    assert list(first[0][len(STUDENT_HEADERS) :]) == [
        "Áo Blazer (Vest) - cỡ đặt",
        "Áo Blazer (Vest) - số lượng đặt",
        "Áo Gile - cỡ đặt",
        "Áo Gile - số lượng đặt",
        *tail,
    ]
    assert list(other[0][len(STUDENT_HEADERS) :]) == [
        "Áo Blazer (Vest) - cỡ đặt",
        "Áo Blazer (Vest) - số lượng đặt",
        *tail,
    ]
    assert list(empty[0][len(STUDENT_HEADERS) :]) == tail
    assert [r[3] for r in first[1:]] == ["Nguyễn Minh Khang"]
    assert [r[3] for r in other[1:]] == ["Lê Thị Ánh"]
    assert len(empty) == 1
    assert not any("suy ra" in str(c) for ws in (first, other, empty) for c in ws[0])


async def test_students_export_school_filter(client: Any, two_schools: dict[str, Any]) -> None:
    other = two_schools["other"]

    csv_rows = read_csv((await client.get(f"/api/export/students.csv?school_id={other}")).content)
    workbook = load_workbook(
        io.BytesIO((await client.get(f"/api/export/students.xlsx?school_id={other}")).content)
    )
    all_xlsx = load_workbook(io.BytesIO((await client.get("/api/export/all.xlsx")).content))

    assert [r[3] for r in csv_rows[1:]] == ["Lê Thị Ánh"]
    assert csv_rows[0][len(STUDENT_HEADERS) :] == [
        "Áo Blazer (Vest) - cỡ đặt",
        "Áo Blazer (Vest) - số lượng đặt",
        "Tạm tính đơn (đ)",
        "Số lần thử",
    ]
    assert workbook.sheetnames == ["Trường Q   LQD  Một"]
    assert all_xlsx.sheetnames == [
        "Trường A",
        "Trường Q   LQD  Một",
        "Trường Trống",
        "Lần thử",
    ]
    missing = await client.get(f"/api/export/students.csv?school_id={uuid.uuid4()}")
    assert missing.status_code == 404


async def test_students_csv_without_filter_uses_union_of_profiles(
    client: Any, two_schools: dict[str, Any], open_campaign: Any
) -> None:
    await open_campaign(two_schools["first"])

    rows = read_csv((await client.get("/api/export/students.csv")).content)

    assert rows[0][len(STUDENT_HEADERS) :] == [
        "Áo Blazer (Vest) - cỡ đặt",
        "Áo Blazer (Vest) - số lượng đặt",
        "Áo Gile - cỡ đặt",
        "Áo Gile - số lượng đặt",
        "Tạm tính đơn (đ)",
        "Số lần thử",
    ]
    assert len(rows) == 3


def test_sheet_titles_are_excel_safe_and_unique() -> None:
    taken: set[str] = set()
    first = _sheet_title("x" * 40, taken)
    second = _sheet_title("x" * 40, taken)
    assert len(first) == 31 and len(second) <= 31 and first != second
    assert _sheet_title("[]:*?/\\", taken) == "Trường"


async def test_campaign_orders_total_is_blank_when_any_line_is_unpriced(
    client: Any, make_student: Any, session: Any, products: Any, open_campaign: Any
) -> None:
    priced = await make_student("Nguyễn Minh Khang")
    mixed = await make_student("Lê Tự Khai")
    campaign = await open_campaign(priced.school_id)
    await place_order(session, campaign, priced, [(products["blazer"], "3", 2, 500_000)])
    await place_order(
        session,
        campaign,
        mixed,
        [(products["blazer"], "3", 1, 500_000), (products["gile"], "2", 1, None)],
    )

    r = await client.get(
        "/api/export/campaign-orders.xlsx", params={"campaign_id": str(campaign.id)}
    )

    headers, *rows = [
        list(x)
        for x in load_workbook(io.BytesIO(r.content))["Bảng kê đơn"].iter_rows(values_only=True)
    ]
    total = headers.index("Thành tiền")
    assert {row[1]: row[total] for row in rows} == {
        "Nguyễn Minh Khang": 1_000_000,
        "Lê Tự Khai": None,
    }


async def test_students_export_ignores_draft_orders(
    client: Any, make_student: Any, session: Any, products: Any, open_campaign: Any
) -> None:
    student = await make_student("Nguyễn Minh Khang")
    campaign = await open_campaign(student.school_id)
    await place_order(
        session, campaign, student, [(products["blazer"], "3", 1, 500_000)], status="DRAFT"
    )

    url = "/api/export/students.csv"
    header, row = student_rows(url, (await client.get(url)).content)
    assert row[header.index("Tạm tính đơn (đ)")] == ""
    assert row[header.index("Áo Blazer (Vest) - cỡ đặt")] == ""
