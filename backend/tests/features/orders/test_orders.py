import io
import uuid
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import SessionLocal
from app.core.errors import ApiError
from app.features.campaigns.models import Campaign
from app.features.catalog.models import MeasurementType, Product
from app.features.orders import handler, lines
from app.features.orders.models import STAFF
from app.features.orders.schemas import LinesIn
from tests.features.campaigns.test_campaigns import (
    item,
    kit,
    make_campaign,
    make_product,
    make_school,
    put_items,
    put_kits,
)

SIZES = ["1", "2", "3", "4", "5"]


class Setup:
    def __init__(self, client: AsyncClient) -> None:
        self.client = client
        self.cid = ""
        self.kit_id = ""
        self.slots: list[str] = []
        self.shirt: dict[str, Any] = {}
        self.skirt: dict[str, Any] = {}
        self.pants: dict[str, Any] = {}
        self.jacket: dict[str, Any] = {}

    async def build(self, jacket_price: int | None = 300_000, status: str = "OPEN") -> "Setup":
        c = self.client
        school = await make_school(c, "TM")
        kw = {"sizes": SIZES}
        self.shirt = await make_product(c, **kw)
        self.skirt = await make_product(c, gender_target="FEMALE", **kw)
        self.pants = await make_product(c, gender_target="MALE", **kw)
        self.jacket = await make_product(c, **kw)
        self.cid = await make_campaign(c, school)
        await put_items(
            c,
            self.cid,
            item(self.shirt, 100_000),
            item(self.skirt, 150_000),
            item(self.pants, 140_000),
            item(self.jacket, jacket_price),
        )
        ids = [self.shirt["id"], self.skirt["id"], self.pants["id"]]
        r = await put_kits(
            c,
            self.cid,
            kit("bo_he", ("Sơ mi", 1, [ids[0]]), ("Dưới", 1, [ids[1], ids[2]])),
        )
        kit_out = r.json()["kits"][0]
        self.kit_id = kit_out["id"]
        self.slots = [s["id"] for s in kit_out["slots"]]
        if status != "DRAFT":
            assert (await c.post(f"/api/campaigns/{self.cid}/open")).status_code == 200
        if status == "LOCKED":
            assert (await c.post(f"/api/campaigns/{self.cid}/lock")).status_code == 200
        return self

    async def order(self, session: AsyncSession, **overrides: Any) -> uuid.UUID:
        campaign = await session.get(Campaign, uuid.UUID(self.cid))
        assert campaign is not None
        fields = {"recipient_name": "Nguyễn Văn A", "gender": "FEMALE", **overrides}
        order = await lines.create_order(session, campaign, source=STAFF, actor="admin", **fields)
        await session.commit()
        return order.id

    def lines_body(self, **overrides: Any) -> dict[str, Any]:
        body: dict[str, Any] = {
            "kit_lines": [{"kit_id": self.kit_id, "qty": 1, "size_label": "3"}],
            "item_lines": [
                {"product_id": self.jacket["id"], "size_label": "4", "qty": 1},
            ],
        }
        body.update(overrides)
        return body


async def put(client: AsyncClient, order_id: uuid.UUID, body: dict[str, Any]) -> Any:
    return await client.put(f"/api/orders/{order_id}/lines", json=body)


async def test_kit_expansion_price_and_total(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    r = await put(client, oid, s.lines_body(submit=True))
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["status"] == "SUBMITTED"
    assert len(out["kit_lines"]) == 1
    by_product = {line["product_id"]: line for line in out["lines"]}
    assert set(by_product) == {s.shirt["id"], s.skirt["id"], s.jacket["id"]}
    assert by_product[s.skirt["id"]]["size_label"] == "3"
    assert by_product[s.skirt["id"]]["kit_line_id"] is not None
    assert by_product[s.jacket["id"]]["kit_line_id"] is None
    assert out["total"] == 100_000 + 150_000 + 300_000
    assert out["item_count"] == 3
    assert out["has_unpriced"] is False


async def test_replace_lines_on_order_created_in_same_session(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    campaign = await session.get(Campaign, uuid.UUID(s.cid))
    assert campaign is not None
    order = await lines.create_order(
        session, campaign, recipient_name="Lê B", source=STAFF, actor="parent", gender="MALE"
    )
    body = LinesIn(**s.lines_body())
    await lines.replace_order_lines(
        session, order, body.kit_lines, body.item_lines, "parent", submit=True, require_open=True
    )
    assert order.status == "SUBMITTED" and len(order.lines) == 3


async def test_kit_qty_multiplies_slot_qty(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    body = s.lines_body(item_lines=[])
    body["kit_lines"][0]["qty"] = 3
    out = (await put(client, oid, body)).json()
    assert [line["qty"] for line in out["lines"]] == [3, 3]


async def test_default_option_follows_gender(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    boy = await s.order(session, gender="MALE", code="GBTM260101_90001")
    out = (await put(client, boy, s.lines_body(item_lines=[]))).json()
    assert s.pants["id"] in {line["product_id"] for line in out["lines"]}


async def test_slot_option_outside_slot_is_422(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    body = s.lines_body()
    body["kit_lines"][0]["slots"] = [{"slot_id": s.slots[0], "product_id": s.skirt["id"]}]
    r = await put(client, oid, body)
    assert r.status_code == 422
    assert "product_id" in "".join(r.json()["error"]["fields"])


@pytest.mark.parametrize(
    ("patch", "field"),
    [
        ({"size_label": "99"}, "size_label"),
        ({"colour": "tím"}, "colour"),
    ],
)
async def test_size_or_colour_outside_product_is_422(
    client: AsyncClient, session: AsyncSession, patch: dict[str, str], field: str
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    line = {"product_id": s.jacket["id"], "size_label": "4", "qty": 1, **patch}
    r = await put(client, oid, s.lines_body(item_lines=[line]))
    assert r.status_code == 422
    assert any(key.endswith(field) for key in r.json()["error"]["fields"])


async def test_kit_without_size_is_422(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    body = s.lines_body(item_lines=[])
    del body["kit_lines"][0]["size_label"]
    assert (await put(client, oid, body)).status_code == 422


async def test_order_default_size_fills_kit_lines(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session, default_size="2")
    body = s.lines_body(item_lines=[])
    del body["kit_lines"][0]["size_label"]
    out = (await put(client, oid, body)).json()
    assert {line["size_label"] for line in out["lines"]} == {"2"}


async def test_price_is_snapshot_and_survives_price_change(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    await put(client, oid, s.lines_body(submit=True))
    await put_items(
        client,
        s.cid,
        item(s.shirt, 999_000),
        item(s.skirt, 150_000),
        item(s.pants, 140_000),
        item(s.jacket, 999_000),
    )
    out = (await client.get(f"/api/orders/{oid}")).json()
    assert out["total"] == 550_000
    body = s.lines_body()
    body["kit_lines"][0]["slots"] = [{"slot_id": s.slots[0], "size_label": "5"}]
    out = (await put(client, oid, body)).json()
    sizes = {line["product_id"]: line["size_label"] for line in out["lines"]}
    assert sizes[s.shirt["id"]] == "5" and sizes[s.skirt["id"]] == "3"
    assert len(out["kit_lines"]) == 1
    assert out["total"] == 550_000


async def test_total_with_unpriced_line(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build(jacket_price=None)
    oid = await s.order(session)
    out = (await put(client, oid, s.lines_body(submit=True))).json()
    assert out["total"] == 250_000
    assert out["has_unpriced"] is True
    listing = (await client.get(f"/api/campaigns/{s.cid}/orders")).json()
    assert listing["items"][0]["has_unpriced"] is True and listing["items"][0]["total"] == 250_000


async def test_one_order_per_student_per_campaign(
    client: AsyncClient, session: AsyncSession, make_student: Any
) -> None:
    s = await Setup(client).build()
    student = await make_student()
    await s.order(session, student_id=student.id)
    with pytest.raises(ApiError) as exc:
        await s.order(session, student_id=student.id)
    assert exc.value.status == 409 and exc.value.code == "order_exists"
    await s.order(session)
    await s.order(session)


async def test_submit_needs_a_line(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    r = await put(client, oid, {"kit_lines": [], "item_lines": [], "submit": True})
    assert r.status_code == 422


async def test_admin_edits_locked_campaign_parent_cannot(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    await put(client, oid, s.lines_body(submit=True))
    assert (await client.post(f"/api/campaigns/{s.cid}/lock")).status_code == 200
    body = s.lines_body()
    body["item_lines"][0]["size_label"] = "5"
    r = await put(client, oid, body)
    assert r.status_code == 200 and r.json()["campaign_status"] == "LOCKED"
    order = await handler.get_order(session, oid)
    with pytest.raises(ApiError) as exc:
        await lines.replace_order_lines(
            session,
            order,
            LinesIn(**body).kit_lines,
            LinesIn(**body).item_lines,
            "parent",
            require_open=True,
        )
    assert exc.value.status == 409 and exc.value.code == "campaign_locked"


async def test_audit_for_every_change_and_publish(
    client: AsyncClient, session: AsyncSession
) -> None:
    from app.core.events import bus

    s = await Setup(client).build()
    async with bus.subscribe() as queue:
        oid = await s.order(session)
        await put(client, oid, s.lines_body(submit=True))
        r = await client.patch(f"/api/orders/{oid}", json={"note": "Áo khoác S5", "gender": None})
        assert r.status_code == 200 and r.json()["note"] == "Áo khoác S5"
        assert r.json()["gender"] is None
        assert queue.get_nowait() == {"orders"}
    rows = await session.execute(
        text("SELECT action FROM audit_logs WHERE entity = 'order' ORDER BY at, id")
    )
    assert [x[0] for x in rows] == ["order.create", "order.lines.replace", "order.update"]
    history = (await client.get(f"/api/orders/{oid}")).json()["history"]
    assert [h["action"] for h in history] == ["order.update", "order.lines.replace", "order.create"]


async def test_list_filters_and_search(client: AsyncClient, session: AsyncSession) -> None:
    s = await Setup(client).build()
    a = await s.order(session, recipient_name="Lê Thị Ánh", class_label="2A1")
    await s.order(session, recipient_name="Trần Bình", gender=None)
    await put(client, a, s.lines_body(submit=True))
    base = f"/api/campaigns/{s.cid}/orders"
    assert (await client.get(base)).json()["total"] == 2
    assert (await client.get(base, params={"status": "SUBMITTED"})).json()["total"] == 1
    assert (await client.get(base, params={"gender": "UNKNOWN"})).json()["total"] == 1
    found = (await client.get(base, params={"q": "anh"})).json()
    assert found["total"] == 1 and found["items"][0]["class_label"] == "2A1"
    code = found["items"][0]["code"]
    assert code.startswith("GBTM") and code[10] == "_" and len(code) == 16
    assert (await client.get(base, params={"q": code.lower()})).json()["total"] == 1


async def test_endpoints_require_admin(anon_client: AsyncClient) -> None:
    ids = uuid.uuid4()
    assert (await anon_client.get(f"/api/orders/{ids}")).status_code == 401
    assert (await anon_client.get(f"/api/campaigns/{ids}/orders")).status_code == 401


async def test_replace_lines_holds_share_lock_on_campaign(
    client: AsyncClient, session: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    s = await Setup(client).build()
    oid = await s.order(session)
    order = await handler.get_order(session, oid)
    body = LinesIn(**s.lines_body())
    probes: list[str] = []
    real_commit = session.commit

    async def probe_then_commit() -> None:
        async with SessionLocal() as other:
            try:
                await other.execute(
                    text("SELECT 1 FROM campaigns WHERE id = :id FOR UPDATE NOWAIT"),
                    {"id": s.cid},
                )
                probes.append("acquired")
            except DBAPIError:
                probes.append("blocked")
        await real_commit()

    monkeypatch.setattr(session, "commit", probe_then_commit)
    await lines.replace_order_lines(session, order, body.kit_lines, body.item_lines, "admin")
    assert probes == ["blocked"]


async def test_audit_snapshots_carry_no_personal_data(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    oid = await s.order(
        session,
        recipient_name="Nguyễn Thị Riêng",
        parent_name="Phụ Huynh",
        parent_phone="0901234567",
    )
    await put(client, oid, s.lines_body(submit=True))
    await client.patch(f"/api/orders/{oid}", json={"class_label": "3A"})
    rows = await session.execute(
        text("SELECT before::text, after::text FROM audit_logs WHERE entity = 'order'")
    )
    dump = " ".join(part or "" for row in rows for part in row)
    for secret in ("Riêng", "Phụ Huynh", "0901234567"):
        assert secret not in dump


async def test_tailor_snapshot_kit_export_and_validation(
    client: AsyncClient, session: AsyncSession
) -> None:
    s = await Setup(client).build()
    definition = {
        "key": "chest",
        "label": "Vòng ngực",
        "instruction": "Đo quanh ngực",
        "required": True,
    }
    for product in (s.shirt, s.jacket):
        row = await session.get(Product, uuid.UUID(product["id"]))
        assert row is not None
        row.tailor_measurements = [definition]
    await session.commit()
    campaign = await session.get(Campaign, uuid.UUID(s.cid))
    assert campaign is not None
    assert all(not item.tailor_enabled for item in campaign.items)
    oid = await s.order(session)
    item_line = {
        "product_id": s.jacket["id"],
        "qty": 2,
        "is_tailor": True,
        "measurements": {"chest": 82.5},
        "note": "Nới rộng",
    }
    body = s.lines_body(item_lines=[item_line], submit=True)
    body["kit_lines"][0]["qty"] = 2
    body["kit_lines"][0]["slots"] = [
        {
            "slot_id": s.slots[0],
            "is_tailor": True,
            "measurements": {"chest": 75},
            "note": "Áo bộ",
        }
    ]
    response = await put(client, oid, body)
    assert response.status_code == 200, response.text
    result = response.json()
    tailor = [line for line in result["lines"] if line["is_tailor"]]
    assert len(tailor) == 2
    assert all(line["qty"] == 2 and line["size_label"] == "May đo" for line in tailor)
    assert {line["note"] for line in tailor} == {"Nới rộng", "Áo bộ"}
    assert all(line["product_tailor_measurements"] == [definition] for line in tailor)
    assert tailor[0]["tailor_measurements"][0]["label"] == "Vòng ngực"
    session.add(MeasurementType(key="chest", label="Vòng ngực", normalized_label="vòng ngực"))
    await session.commit()
    renamed = await client.patch("/api/measurement-types/chest", json={"label": "Chu vi ngực"})
    assert renamed.status_code == 200
    line_labels = await session.scalars(
        text("SELECT tailor_measurements->0->>'label' FROM order_lines WHERE is_tailor")
    )
    assert set(line_labels) == {"Vòng ngực"}
    product_labels = await session.scalars(
        text(
            "SELECT tailor_measurements->0->>'label' FROM products WHERE id = ANY(:ids)"
        ).bindparams(ids=[uuid.UUID(s.shirt["id"]), uuid.UUID(s.jacket["id"])])
    )
    assert set(product_labels) == {"Chu vi ngực"}
    response = await client.get(f"/api/export/campaign-orders.xlsx?campaign_id={s.cid}")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    exported = list(workbook["May đo"].iter_rows(min_row=3, values_only=True))
    assert len(exported) == 2
    assert any("Vòng ngực: 82.5" in row[6] and row[7] == "Nới rộng" for row in exported)
    pivot = (await client.get(f"/api/stats/campaign-pivot?campaign_id={s.cid}&group=kit")).json()
    assert any(
        row["size_label"] == "May đo" and row["total"] == 2
        for block in pivot["blocks"]
        for row in block["rows"]
    )

    jacket = await session.get(Product, uuid.UUID(s.jacket["id"]))
    assert jacket is not None
    jacket.tailor_measurements = []
    await session.commit()
    saved = next(line for line in tailor if line["product_id"] == s.jacket["id"])
    item_line["line_id"] = saved["id"]
    body["kit_lines"][0]["slots"][0]["line_id"] = next(
        line["id"] for line in tailor if line["product_id"] == s.shirt["id"]
    )
    response = await put(client, oid, body)
    assert response.status_code == 200, response.text
    saved = next(line for line in response.json()["lines"] if line["product_id"] == s.jacket["id"])
    assert saved["product_tailor_measurements"] == []
    assert saved["tailor_measurements"] == [{**definition, "value": 82.5}]
    item_line["line_id"] = saved["id"]
    body["kit_lines"][0]["slots"][0]["line_id"] = next(
        line["id"] for line in response.json()["lines"] if line["product_id"] == s.shirt["id"]
    )
    assert any(
        h["after"] and any(line.get("tailor_measurements") for line in h["after"].get("lines", []))
        for h in response.json()["history"]
    )
    item_line["measurements"] = {"chest": 0}
    assert (await put(client, oid, body)).status_code == 422
    item_line["measurements"] = {}
    assert (await put(client, oid, body)).status_code == 422
    item_line["measurements"] = {"chest": 82}
    item_line.pop("line_id")
    assert (await put(client, oid, body)).status_code == 422
