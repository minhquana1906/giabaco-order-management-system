from pathlib import Path
from typing import Any

import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.audit import AuditLog
from app.features.campaigns.models import LOCKED, Campaign
from app.features.orders import legacy_thtm as legacy
from app.features.orders.models import DRAFT, LEGACY_IMPORT, SUBMITTED, Order
from app.features.students.models import Student

SKU_COLS = {s.code: 10 + i for i, s in enumerate(legacy.SKUS)}
KIT_COLS = {k.code: 20 + i for i, k in enumerate(legacy.KIT_SPECS)}
PRICES = [105000, 115000, 105000, 100000, 100000, 105000, 100000, 100000, 105000, 195000]


def counts(**kw: int) -> dict[str, int]:
    return kw


def kit_qty(plan: legacy.KitPlan) -> dict[str, int]:
    return {u.kit: u.qty for u in plan.kits}


def test_boy_buys_shorts_kit() -> None:
    plan = legacy.assign_kits(counts(so_mi_coc=2, sooc_kaki=2), "MALE")
    assert plan.kits == (legacy.KitUse("bo_he", 2, "sooc_kaki"),)
    assert plan.loose == {}


def test_girl_buys_skirt_kit() -> None:
    plan = legacy.assign_kits(counts(so_mi_coc=1, vay_kaki=1), "FEMALE")
    assert plan.kits == (legacy.KitUse("bo_he", 1, "vay_kaki"),)


def test_more_shirts_than_bottoms_leaves_loose_shirts() -> None:
    plan = legacy.assign_kits(counts(so_mi_coc=3, sooc_kaki=1, so_mi_dai=1), "MALE")
    assert kit_qty(plan) == {"bo_he": 1}
    assert plan.loose == {"so_mi_coc": 2, "so_mi_dai": 1}


def test_more_bottoms_than_shirts_leaves_loose_bottoms() -> None:
    plan = legacy.assign_kits(counts(so_mi_dai=1, quan_dai_kaki=3, vay_kaki=2), None)
    assert kit_qty(plan) == {"bo_dai": 1}
    assert plan.loose == {"quan_dai_kaki": 2, "vay_kaki": 2}


def test_sports_kits_follow_excel_formulas() -> None:
    plan = legacy.assign_kits(counts(ao_phong_tt=2, sooc_tt=1, quan_dai_tt=1, ao_khoac=1), "MALE")
    assert kit_qty(plan) == {"tt_coc": 1, "tt_dai": 1}
    assert plan.loose == {"ao_khoac": 1}


def test_sports_long_kit_never_negative() -> None:
    plan = legacy.assign_kits(counts(quan_dai_tt=2), None)
    assert plan.kits == ()
    assert plan.loose == {"quan_dai_tt": 2}


def test_mixed_bottoms_use_larger_group_and_keep_the_rest_loose() -> None:
    plan = legacy.assign_kits(counts(so_mi_coc=2, sooc_kaki=1, vay_kaki=1), "FEMALE")
    assert plan.kits == (legacy.KitUse("bo_he", 1, "vay_kaki"),)
    assert plan.loose == {"so_mi_coc": 1, "sooc_kaki": 1}


def test_loose_only_order() -> None:
    plan = legacy.assign_kits(counts(ao_khoac=2), "MALE")
    assert plan.kits == ()
    assert plan.loose == {"ao_khoac": 2}


def build_sheet(path: Path, header_shift: str | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = legacy.SHEET
    top = {5: "Mã KH", 6: "Họ tên Học sinh", 7: "Họ tên Phụ huynh", 8: "Giới tính", 9: "Size"}
    top |= {34: "Tổng số sp", 35: "Ghi chú"}
    for col, text in top.items():
        ws.cell(8, col, text)
    ws.cell(9, 4, "Check kiểm")
    for sku in legacy.SKUS:
        ws.cell(
            9, SKU_COLS[sku.code], sku.header.upper() + (" x" if sku.code == header_shift else "")
        )
    for spec in legacy.KIT_SPECS:
        ws.cell(9, KIT_COLS[spec.code], spec.name + " (mô tả)")
    ws.cell(10, 3, "Đơn giá")
    for sku, price in zip(legacy.SKUS, PRICES, strict=True):
        ws.cell(10, SKU_COLS[sku.code], price)
    ws.cell(11, 1, "Bao 1")
    boy = {"so_mi_coc": 2, "sooc_kaki": 1, "so_mi_dai": 1, "quan_dai_kaki": 1, "ao_khoac": 1}
    orders: list[tuple[str, str, Any, dict[str, int], str | None, str | None]] = [
        ("GBTM260701_10001", "Nam", 3, boy, "Thiếu Áo khoác", "Ao khoac S5"),
        ("GBTM260702_000002", "Nữ", 2, {"so_mi_coc": 1, "vay_kaki": 1}, None, None),
        ("GBTM260702_000003", "Trống", None, {}, None, None),
    ]  # fmt: skip
    for n, (code, gender, size, items, check, note) in enumerate(orders):
        r = 12 + n
        ws.cell(r, 1, n + 1)
        ws.cell(r, 4, check)
        ws.cell(r, 5, code)
        ws.cell(r, 6, "Học sinh Một")
        ws.cell(r, 7, "Phụ huynh Một")
        ws.cell(r, 8, gender)
        ws.cell(r, 9, size)
        ws.cell(r, 35, note)
        for code_, col in SKU_COLS.items():
            ws.cell(r, col, items.get(code_, 0))
    ws.cell(15, 7, "Tổng")
    ws.cell(15, 34, 8)
    for code, qty in {"bo_he": 2, "bo_dai": 1, "tt_coc": 0, "tt_dai": "0"}.items():
        ws.cell(15, KIT_COLS[code], qty)
    ws.cell(16, 5, "Tổng cộng")
    ws.cell(16, 34, 8)
    for code, qty in {"bo_he": 2, "bo_dai": 1, "tt_coc": 0, "tt_dai": 0}.items():
        ws.cell(16, KIT_COLS[code], qty)
    wb.save(path)
    return path


def test_parser_reads_by_header_and_keeps_names_out_of_repr(tmp_path: Path) -> None:
    sheet = legacy.parse_sheet(build_sheet(tmp_path / "t.xlsx"))
    assert [r.code for r in sheet.rows] == [
        "GBTM260701_10001",
        "GBTM260702_000002",
        "GBTM260702_000003",
    ]
    first, girl, empty = sheet.rows
    assert (first.gender, first.size, first.bag) == ("MALE", "3", 1)
    assert first.note == "Bao 1 · Thiếu Áo khoác · Ao khoac S5"
    assert first.note_mentions_size and not girl.note_mentions_size
    assert girl.gender == "FEMALE" and girl.note == "Bao 1"
    assert (empty.gender, empty.size, empty.counts) == (None, None, {})
    assert str(first.submitted_on) == "2026-07-01"
    assert sheet.prices["ao_khoac"] == 195000
    assert "Học sinh" not in repr(first)
    assert [t.bag for t in sheet.bag_totals] == [1, 1]


def test_bag_reconciliation_on_synthetic_sheet(tmp_path: Path) -> None:
    sheet = legacy.parse_sheet(build_sheet(tmp_path / "t.xlsx"))
    checks = [
        c for c in legacy.reconcile(sheet, legacy.build_plans(sheet)) if c.label.startswith("bag")
    ]
    assert checks and all(c.ok for c in checks)


def test_changed_header_stops_instead_of_guessing(tmp_path: Path) -> None:
    with pytest.raises(legacy.LegacyError, match="header differs"):
        legacy.parse_sheet(build_sheet(tmp_path / "t.xlsx", header_shift="ao_khoac"))


@pytest.mark.asyncio
async def test_load_writes_orders_with_slots_and_audit(
    session: AsyncSession, tmp_path: Path
) -> None:
    sheet = legacy.parse_sheet(build_sheet(tmp_path / "t.xlsx"))
    await legacy.load(session, sheet, legacy.build_plans(sheet))
    await session.commit()
    campaign = await session.scalar(select(Campaign).where(Campaign.name == legacy.CAMPAIGN_NAME))
    assert campaign is not None and campaign.status == LOCKED
    assert len(campaign.items) == 10 and len(campaign.kits) == 4
    orders = {o.code: o for o in (await session.scalars(select(Order))).all()}
    first = orders["GBTM260701_10001"]
    assert (first.status, first.source, first.student_id) == (SUBMITTED, LEGACY_IMPORT, None)
    assert first.submitted_at is not None and first.submitted_at.date().isoformat() == "2026-07-01"
    assert {line.kit_line_id is not None for line in first.lines} == {True, False}
    assert all(line.slot_id for line in first.lines if line.kit_line_id)
    assert all(line.size_label == "3" for line in first.lines)
    assert (
        sum(line.qty * (line.unit_price or 0) for line in first.lines)
        == 2 * 105000 + 100000 + 115000 + 105000 + 195000
    )
    draft = orders["GBTM260702_000003"]
    assert (draft.status, draft.lines, draft.submitted_at) == (DRAFT, [], None)
    actors = (await session.scalars(select(AuditLog.actor).distinct())).all()
    assert set(actors) == {legacy.ACTOR}
    assert await session.scalar(select(func.count()).select_from(Student)) == 0


@pytest.mark.asyncio
async def test_load_twice_stops_without_writing(session: AsyncSession, tmp_path: Path) -> None:
    sheet = legacy.parse_sheet(build_sheet(tmp_path / "t.xlsx"))
    plans = legacy.build_plans(sheet)
    await legacy.load(session, sheet, plans)
    await session.commit()
    with pytest.raises(legacy.LegacyError, match="already exist"):
        await legacy.load(session, sheet, plans)
    assert await session.scalar(select(func.count()).select_from(Order)) == 3
