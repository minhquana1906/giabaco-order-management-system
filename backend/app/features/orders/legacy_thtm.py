"""One-off loader for the Tay Mo (THTM) legacy sheet: legacy_thtm FILE [--dry-run|--commit]."""

import argparse
import asyncio
import re
import sys
from collections import defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import audit
from app.core.db import engine
from app.features.campaigns.models import (
    LOCKED,
    Campaign,
    CampaignItem,
    CampaignKit,
    KitSlot,
    KitSlotOption,
)
from app.features.catalog.models import Product, School
from app.features.orders.lines import (
    ItemLineIn,
    KitLineIn,
    SlotPick,
    create_order,
    replace_order_lines,
)
from app.features.orders.models import DRAFT, LEGACY_IMPORT, Order

ACTOR = "system:legacy-thtm"
SHEET = "TỔNG DS"
SCHOOL_CODE = "THTM"
SCHOOL_NAME = "Trường Tiểu học Tây Mỗ"
CAMPAIGN_NAME = "Tây Mỗ 2026–2027"
SCHOOL_YEAR = "2026-2027"
COLOUR = "ghi"
SIZES = [str(n) for n in range(1, 9)]
SUBMIT_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
GENDERS = {"nam": "MALE", "nữ": "FEMALE", "trống": None}
BAG_LABEL = re.compile(r"(?:bao|bọc)\s*(\d+)")
CODE_DATE = re.compile(r"(\d{6})_")
NOTE_HAS_SIZE = re.compile(r"\d")
MAX_NOTE = 500
MAX_KIT_QTY = 10


class LegacyError(Exception):
    """A failed load step; the message is safe to print (no personal data)."""


@dataclass(frozen=True, slots=True)
class Sku:
    code: str
    header: str
    name: str
    gender_target: str = "UNISEX"


SKUS = (
    Sku("so_mi_coc", "áo trắng sơ mi cộc tay", "Áo sơ mi trắng cộc tay"),
    Sku("so_mi_dai", "áo trắng sơ my dài tay", "Áo sơ mi trắng dài tay"),
    Sku("ao_phong_tt", "áo phông thể thao", "Áo phông thể thao"),
    Sku("sooc_tt", "quần sooc tt 2 sọc", "Quần sooc thể thao 2 sọc"),
    Sku("juyp_tt", "juyp tt 2 sọc", "Juyp thể thao 2 sọc", "FEMALE"),
    Sku("quan_dai_tt", "quần dài tt 2 sọc", "Quần dài thể thao 2 sọc"),
    Sku("vay_kaki", "váy kaki nữ", "Váy kaki nữ", "FEMALE"),
    Sku("sooc_kaki", "quần soóc nam kaki", "Quần sooc kaki nam", "MALE"),
    Sku("quan_dai_kaki", "quần dài kaki tím than", "Quần dài kaki tím than"),
    Sku("ao_khoac", "áo khoác bo đai lót tricot", "Áo khoác bo đai lót tricot"),
)
PRODUCT_CODE_PREFIX = "thtm_"


@dataclass(frozen=True, slots=True)
class KitSpec:
    code: str
    name: str
    header: str
    top: str
    bottoms: tuple[str, ...]


KIT_SPECS = (
    KitSpec("bo_he", "Bộ hè", "bộ hè", "so_mi_coc", ("sooc_kaki", "vay_kaki")),
    KitSpec("bo_dai", "Bộ dài", "bộ dài", "so_mi_dai", ("quan_dai_kaki",)),
    KitSpec("tt_coc", "Bộ thể thao cộc", "bộ thể thao cộc", "ao_phong_tt", ("sooc_tt", "juyp_tt")),
    KitSpec("tt_dai", "Bộ thể thao dài", "bộ thể thao dài", "ao_phong_tt", ("quan_dai_tt",)),
)
KIT_BY_CODE = {k.code: k for k in KIT_SPECS}

EXPECTED_ORDERS = 407
EXPECTED_ITEMS = 3167
EXPECTED_MONEY = 356_780_000
EXPECTED_KITS = {"bo_he": 713, "bo_dai": 108, "tt_coc": 452, "tt_dai": 29}

ROW8_HEADERS = ("mã kh", "họ tên học sinh", "họ tên phụ huynh", "giới tính", "size", "ghi chú")
TOTAL_ITEMS_HEADER = "tổng số sp"
CHECK_HEADER = "check kiểm"


@dataclass(frozen=True, slots=True)
class KitUse:
    kit: str
    qty: int
    bottom: str


@dataclass(frozen=True, slots=True)
class KitPlan:
    kits: tuple[KitUse, ...]
    loose: dict[str, int]


def assign_kits(counts: Mapping[str, int], gender: str | None) -> KitPlan:
    """Split per-SKU counts into kits (Excel formulas T..W) and loose items."""
    c = defaultdict(int, counts)
    qty = {
        "bo_he": min(c["so_mi_coc"], max(c["sooc_kaki"], c["vay_kaki"])),
        "bo_dai": min(c["quan_dai_kaki"], c["so_mi_dai"]),
        "tt_coc": min(c["ao_phong_tt"], max(c["sooc_tt"], c["juyp_tt"])),
        "tt_dai": min(c["quan_dai_tt"], c["ao_phong_tt"] - c["sooc_tt"] - c["juyp_tt"]),
    }
    left = defaultdict(int, counts)
    uses: list[KitUse] = []
    for spec in KIT_SPECS:
        n = max(0, qty[spec.code])
        if n == 0:
            continue
        bottom = max(
            reversed(spec.bottoms) if gender == "FEMALE" else spec.bottoms, key=lambda b: left[b]
        )
        left[spec.top] -= n
        left[bottom] -= n
        uses.append(KitUse(spec.code, n, bottom))
    if any(v < 0 for v in left.values()):
        raise LegacyError("kit assignment consumes more items than the order has")
    return KitPlan(tuple(uses), {k: v for k, v in left.items() if v})


@dataclass(frozen=True, slots=True)
class LegacyRow:
    code: str
    bag: int
    gender: str | None
    size: str | None
    counts: dict[str, int]
    note: str | None
    submitted_on: date
    note_mentions_size: bool
    recipient_name: str = field(repr=False)
    parent_name: str | None = field(repr=False)


@dataclass(frozen=True, slots=True)
class BagTotal:
    bag: int
    items: int
    kits: dict[str, int]


@dataclass(frozen=True, slots=True)
class Sheet:
    prices: dict[str, int]
    rows: list[LegacyRow]
    bag_totals: list[BagTotal]


def _norm(value: object) -> str:
    return " ".join(str(value).split()).casefold()


def _int(value: object, what: str) -> int:
    if value is None or value == "":
        return 0
    try:
        number = float(str(value)) if isinstance(value, str) else float(value)  # type: ignore[arg-type]
    except ValueError:
        raise LegacyError(f"non-numeric value in {what}") from None
    if number != int(number) or number < 0:
        raise LegacyError(f"invalid count in {what}")
    return int(number)


def _first_columns(row: Sequence[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(row):
        if isinstance(value, str):
            found.setdefault(_norm(value), index)
    return found


def _require(columns: Mapping[str, int], names: Sequence[str]) -> None:
    missing = [n for n in names if n not in columns]
    if missing:
        raise LegacyError(f"header differs from the verified sheet, missing: {', '.join(missing)}")


def parse_sheet(path: Path) -> Sheet:
    """Read the sheet by header names; stop instead of guessing when the layout differs."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in workbook.sheetnames:
        raise LegacyError(f"sheet {SHEET} not found")
    grid = list(workbook[SHEET].iter_rows(values_only=True))
    head = next((i for i, row in enumerate(grid) if "mã kh" in _first_columns(row)), None)
    if head is None or head + 2 >= len(grid):
        raise LegacyError("header row not found")
    top, sub, price_row = _first_columns(grid[head]), _first_columns(grid[head + 1]), grid[head + 2]
    _require(top, (*ROW8_HEADERS, TOTAL_ITEMS_HEADER))
    _require(sub, (CHECK_HEADER, *(s.header for s in SKUS)))
    kit_cols = {}
    for spec in KIT_SPECS:
        col = next((i for n, i in sub.items() if n.startswith(spec.header)), None)
        if col is None:
            raise LegacyError(f"header differs from the verified sheet, missing: {spec.header}")
        kit_cols[spec.code] = col
    sku_cols = {s.code: sub[s.header] for s in SKUS}
    prices = {code: _int(price_row[col], f"price {code}") for code, col in sku_cols.items()}
    if not all(prices.values()):
        raise LegacyError("unit price row is missing a price")

    c_code, c_name, c_parent = top["mã kh"], top["họ tên học sinh"], top["họ tên phụ huynh"]
    c_gender, c_size, c_note = top["giới tính"], top["size"], top["ghi chú"]
    c_items, c_check = top[TOTAL_ITEMS_HEADER], sub[CHECK_HEADER]
    rows: list[LegacyRow] = []
    totals: list[BagTotal] = []
    bag = 0
    for row in grid[head + 3 :]:
        code = row[c_code]
        label = _norm(row[0]) if isinstance(row[0], str) else ""
        if isinstance(code, str) and code.startswith("GB"):
            rows.append(
                _order_row(
                    row, bag, sku_cols, c_code, c_name, c_parent, c_gender, c_size, c_note, c_check
                )
            )
        elif m := BAG_LABEL.fullmatch(label):
            bag = int(m.group(1))
        elif any(
            isinstance(v, str) and _norm(v) in ("tổng", "tổng cộng") for v in row[: c_code + 3]
        ):
            totals.append(
                BagTotal(
                    bag,
                    _int(row[c_items], "bag total"),
                    {k: _int(row[c], f"bag {bag} total {k}") for k, c in kit_cols.items()},
                )
            )
            if any(isinstance(v, str) and _norm(v) == "tổng cộng" for v in row[: c_code + 3]):
                break
    return Sheet(prices, rows, totals)


def _order_row(
    row: Sequence[object], bag: int, sku_cols: Mapping[str, int],
    c_code: int, c_name: int, c_parent: int, c_gender: int, c_size: int, c_note: int, c_check: int,
) -> LegacyRow:  # fmt: skip
    code = str(row[c_code])
    stamp = CODE_DATE.search(code)
    if stamp is None:
        raise LegacyError(f"{code}: no yymmdd date in the customer code")
    gender_text = _norm(row[c_gender]) if row[c_gender] is not None else "trống"
    if gender_text not in GENDERS:
        raise LegacyError(f"{code}: unknown gender value")
    size = None if row[c_size] in (None, "") else str(_int(row[c_size], f"{code} size"))
    counts = {k: n for k, c in sku_cols.items() if (n := _int(row[c], f"{code} {k}"))}
    if counts and size is None:
        raise LegacyError(f"{code}: has items but no size")
    parts = [f"Bao {bag}"] + [
        str(row[c]).strip() for c in (c_check, c_note) if row[c] not in (None, "")
    ]
    note = " · ".join(parts)
    if len(note) > MAX_NOTE:
        raise LegacyError(f"{code}: note longer than {MAX_NOTE} characters")
    yy, mm, dd = (int(stamp.group(1)[i : i + 2]) for i in (0, 2, 4))
    try:
        submitted_on = date(2000 + yy, mm, dd)
    except ValueError:
        raise LegacyError(f"{code}: invalid date in the customer code") from None
    parent = row[c_parent]
    return LegacyRow(
        code=code,
        bag=bag,
        gender=GENDERS[gender_text],
        size=size,
        counts=counts,
        note=note,
        submitted_on=submitted_on,
        note_mentions_size=bool(NOTE_HAS_SIZE.search(str(row[c_note] or ""))),
        recipient_name=str(row[c_name]).strip(),
        parent_name=str(parent).strip() if parent not in (None, "") else None,
    )


@dataclass(frozen=True, slots=True)
class Check:
    label: str
    expected: int
    actual: int

    @property
    def ok(self) -> bool:
        return self.expected == self.actual


def _kit_totals(plans: Sequence[KitPlan]) -> dict[str, int]:
    totals = dict.fromkeys(KIT_BY_CODE, 0)
    for plan in plans:
        for use in plan.kits:
            totals[use.kit] += use.qty
    return totals


def reconcile(sheet: Sheet, plans: Sequence[KitPlan]) -> list[Check]:
    kits = _kit_totals(plans)
    checks = [
        Check("orders", EXPECTED_ORDERS, len(sheet.rows)),
        Check("items", EXPECTED_ITEMS, sum(sum(r.counts.values()) for r in sheet.rows)),
        Check(
            "money",
            EXPECTED_MONEY,
            sum(n * sheet.prices[k] for r in sheet.rows for k, n in r.counts.items()),
        ),
        *(Check(f"kit {code}", EXPECTED_KITS[code], kits[code]) for code in KIT_BY_CODE),
    ]
    for total in sheet.bag_totals:
        in_bag = [(r, p) for r, p in zip(sheet.rows, plans, strict=True) if r.bag == total.bag]
        bag_kits = _kit_totals([p for _, p in in_bag])
        checks.append(
            Check(
                f"bag {total.bag} items",
                total.items,
                sum(sum(r.counts.values()) for r, _ in in_bag),
            )
        )
        checks += [
            Check(f"bag {total.bag} kit {k}", total.kits[k], bag_kits[k]) for k in KIT_BY_CODE
        ]
    if len({r.code for r in sheet.rows}) != len(sheet.rows):
        checks.append(
            Check("duplicate codes", 0, len(sheet.rows) - len({r.code for r in sheet.rows}))
        )
    if {r.bag for r in sheet.rows} != {t.bag for t in sheet.bag_totals}:
        checks.append(
            Check("bags with a total row", len({r.bag for r in sheet.rows}), len(sheet.bag_totals))
        )
    return checks


def build_plans(sheet: Sheet) -> list[KitPlan]:
    plans = [assign_kits(r.counts, r.gender) for r in sheet.rows]
    for row, plan in zip(sheet.rows, plans, strict=True):
        if any(use.qty > MAX_KIT_QTY for use in plan.kits):
            raise LegacyError(f"{row.code}: more than {MAX_KIT_QTY} of one kit")
    return plans


async def _ensure_products(session: AsyncSession) -> dict[str, Product]:
    codes = {PRODUCT_CODE_PREFIX + s.code: s for s in SKUS}
    found = {
        p.code: p for p in (await session.scalars(select(Product).where(Product.code.in_(codes))))
    }
    position = 100
    for code, sku in codes.items():
        if code in found:
            continue
        product = Product(
            code=code,
            name=sku.name,
            sizes=SIZES,
            colours=[COLOUR],
            gender_target=sku.gender_target,
            active=True,
            position=position,
        )
        position += 1
        session.add(product)
        await session.flush()
        await audit.record(
            session, ACTOR, "product.create", "product", product.id, None,
            {"code": code, "name": sku.name, "sizes": SIZES, "colours": [COLOUR]},
        )  # fmt: skip
        found[code] = product
    return {code.removeprefix(PRODUCT_CODE_PREFIX): p for code, p in found.items()}


async def _ensure_campaign(
    session: AsyncSession, prices: Mapping[str, int], products: Mapping[str, Product]
) -> Campaign:
    school = await session.scalar(select(School).where(School.code == SCHOOL_CODE))
    if school is None:
        school = School(name=SCHOOL_NAME, code=SCHOOL_CODE)
        session.add(school)
        await session.flush()
        await audit.record(
            session, ACTOR, "school.create", "school", school.id, None,
            {"name": SCHOOL_NAME, "code": SCHOOL_CODE},
        )  # fmt: skip
    campaign = await session.scalar(
        select(Campaign).where(Campaign.school_id == school.id, Campaign.name == CAMPAIGN_NAME)
    )
    created = campaign is None
    if campaign is None:
        campaign = Campaign(
            school_id=school.id,
            name=CAMPAIGN_NAME,
            school_year=SCHOOL_YEAR,
            opens_at=datetime(2026, 7, 1, tzinfo=SUBMIT_TZ),
            status=LOCKED,
            items=[],
            kits=[],
        )
        session.add(campaign)
    else:
        live = await session.scalar(
            text("SELECT count(*) FROM orders WHERE campaign_id = :c AND deleted_at IS NULL"),
            {"c": campaign.id},
        )
        if live:
            raise LegacyError("campaign already has orders; nothing was written")
        campaign.items.clear()
        await session.flush()
    campaign.items.extend(
        CampaignItem(
            product_id=products[s.code].id, unit_price=prices[s.code], colour=COLOUR, sort=n
        )
        for n, s in enumerate(SKUS)
    )
    existing = {k.code: k for k in campaign.kits}
    stale = set(existing) - set(KIT_BY_CODE)
    if stale:
        raise LegacyError(f"campaign has unexpected kits: {', '.join(sorted(stale))}")
    for n, spec in enumerate(KIT_SPECS):
        slots = [
            KitSlot(
                label=label,
                qty=1,
                sort=i,
                options=[KitSlotOption(product_id=products[c].id) for c in codes],
            )
            for i, (label, codes) in enumerate((("Áo", (spec.top,)), ("Quần/Váy", spec.bottoms)))
        ]
        kit = existing.get(spec.code)
        if kit is None:
            campaign.kits.append(CampaignKit(code=spec.code, name=spec.name, sort=n, slots=slots))
        else:
            kit.name, kit.sort, kit.slots = spec.name, n, slots
    await session.flush()
    campaign.status = LOCKED
    await audit.record(
        session, ACTOR, "campaign.create" if created else "campaign.reshape", "campaign",
        campaign.id, None, {"name": CAMPAIGN_NAME, "status": LOCKED},
    )  # fmt: skip
    return campaign


def _lines(
    plan: KitPlan, row: LegacyRow, campaign: Campaign, products: Mapping[str, Product]
) -> tuple[list[KitLineIn], list[ItemLineIn]]:
    kits = {k.code: k for k in campaign.kits}
    kit_lines = []
    for use in plan.kits:
        spec = KIT_BY_CODE[use.kit]
        top_slot, bottom_slot = kits[use.kit].slots
        kit_lines.append(
            KitLineIn(
                kit_id=kits[use.kit].id,
                qty=use.qty,
                slots=[
                    SlotPick(slot_id=top_slot.id, product_id=products[spec.top].id),
                    SlotPick(slot_id=bottom_slot.id, product_id=products[use.bottom].id),
                ],
            )
        )
    assert row.size is not None or not plan.loose
    item_lines = [
        ItemLineIn(product_id=products[k].id, size_label=row.size or "", qty=n)
        for k, n in plan.loose.items()
    ]
    return kit_lines, item_lines


async def load(session: AsyncSession, sheet: Sheet, plans: Sequence[KitPlan]) -> list[Check]:
    """Write everything into the caller's transaction and reconcile against the database."""
    codes = [r.code for r in sheet.rows]
    taken = (await session.scalars(select(Order.code).where(Order.code.in_(codes)))).all()
    if taken:
        raise LegacyError(f"{len(taken)} order codes already exist; nothing was written")
    products = await _ensure_products(session)
    campaign = await _ensure_campaign(session, sheet.prices, products)
    for row, plan in zip(sheet.rows, plans, strict=True):
        order = await create_order(
            session,
            campaign,
            recipient_name=row.recipient_name,
            parent_name=row.parent_name,
            gender=row.gender,
            default_size=row.size,
            note=row.note,
            source=LEGACY_IMPORT,
            actor=ACTOR,
            code=row.code,
        )
        if not row.counts:
            assert order.status == DRAFT
            continue
        kit_lines, item_lines = _lines(plan, row, campaign, products)
        await replace_order_lines(session, order, kit_lines, item_lines, ACTOR, submit=True)
        order.submitted_at = datetime.combine(row.submitted_on, time(12), SUBMIT_TZ)
    await session.flush()
    return await _database_checks(session, campaign, sheet)


async def _database_checks(session: AsyncSession, campaign: Campaign, sheet: Sheet) -> list[Check]:
    params = {"c": campaign.id}
    orders = await session.scalar(
        text("SELECT count(*) FROM orders WHERE campaign_id = :c AND deleted_at IS NULL"), params
    )
    items, money = (
        await session.execute(
            text(
                "SELECT coalesce(sum(l.qty), 0), coalesce(sum(l.qty * l.unit_price), 0) "
                "FROM order_lines l JOIN orders o ON o.id = l.order_id "
                "WHERE o.campaign_id = :c AND o.deleted_at IS NULL"
            ),
            params,
        )
    ).one()
    kit_rows = (
        await session.execute(
            text(
                "SELECT k.code, coalesce(sum(kl.qty) FILTER (WHERE o.id IS NOT NULL), 0) "
                "FROM campaign_kits k "
                "LEFT JOIN order_kit_lines kl ON kl.kit_id = k.id "
                "LEFT JOIN orders o ON o.id = kl.order_id AND o.deleted_at IS NULL "
                "WHERE k.campaign_id = :c GROUP BY k.code"
            ),
            params,
        )
    ).all()
    kits = {code: int(n) for code, n in kit_rows}
    return [
        Check("db orders", EXPECTED_ORDERS, int(orders or 0)),
        Check("db items", EXPECTED_ITEMS, int(items)),
        Check("db money", EXPECTED_MONEY, int(money)),
        *(Check(f"db kit {code}", n, kits.get(code, 0)) for code, n in EXPECTED_KITS.items()),
    ]


def _print_checks(checks: Sequence[Check]) -> None:
    for check in checks:
        mark = "OK  " if check.ok else "FAIL"
        print(f"{mark} {check.label:<22} expected={check.expected:>11,} actual={check.actual:>11,}")


def _print_review(sheet: Sheet) -> None:
    empty = [r.code for r in sheet.rows if not r.counts]
    sized = [r.code for r in sheet.rows if r.note_mentions_size]
    print(f"orders without lines (left DRAFT): {', '.join(empty) or '-'}")
    print(
        f"orders whose note mentions a size, fix by hand ({len(sized)}): {', '.join(sized) or '-'}"
    )


async def _commit(sheet: Sheet, plans: Sequence[KitPlan]) -> list[Check]:
    async with engine.connect() as conn:
        transaction = await conn.begin()
        session = AsyncSession(
            conn, join_transaction_mode="create_savepoint", expire_on_commit=False
        )
        try:
            checks = await load(session, sheet, plans)
            if all(c.ok for c in checks):
                await transaction.commit()
            else:
                await transaction.rollback()
        except BaseException:
            await transaction.rollback()
            raise
        finally:
            await session.close()
    return checks


def main(argv: Sequence[str]) -> int:
    parser = argparse.ArgumentParser(prog="legacy_thtm", description=__doc__)
    parser.add_argument("file", type=Path)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--dry-run", action="store_true", help="print the reconciliation only (default)"
    )
    mode.add_argument("--commit", action="store_true", help="write in one transaction")
    args = parser.parse_args(argv)
    try:
        sheet = parse_sheet(args.file)
        plans = build_plans(sheet)
        checks = reconcile(sheet, plans)
        _print_checks(checks)
        _print_review(sheet)
        if not all(c.ok for c in checks):
            print("Reconciliation failed; nothing was written.")
            return 1
        if not args.commit:
            print("Dry run: nothing was written.")
            return 0
        db_checks = asyncio.run(_commit(sheet, plans))
        _print_checks(db_checks)
        if not all(c.ok for c in db_checks):
            print("Database reconciliation failed; rolled back.")
            return 1
    except LegacyError as error:
        print(f"Stopped: {error}")
        return 1
    print("Committed.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
